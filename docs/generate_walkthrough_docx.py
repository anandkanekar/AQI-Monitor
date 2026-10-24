#!/usr/bin/env python3
"""Generates docs/AQI_Project_Walkthrough.docx - a Word walkthrough of the AQI Tracking System.

Usage:
    pip install python-docx
    python docs/generate_walkthrough_docx.py
"""
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOCS_DIR = Path(__file__).resolve().parent
ROOT = DOCS_DIR.parent
OUT = DOCS_DIR / "AQI_Project_Walkthrough.docx"
ARCH_PNG = DOCS_DIR / "architecture.png"

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
ACCENT = RGBColor(0x1E, 0x88, 0xE5)
GREY = RGBColor(0x54, 0x6E, 0x7A)
LIGHT_FILL = "EEF3F8"
HEADER_FILL = "1B2A4A"


def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = tcPr.makeelement(qn("w:shd"), {})
        tcPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)


def set_cell_text(cell, text, bold=False, color=None, size=10, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def build_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal.paragraph_format.space_after = Pt(6)

    for i, size in [(1, 26), (2, 18), (3, 14), (4, 12)]:
        h = doc.styles[f"Heading {i}"]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = NAVY
        h.paragraph_format.space_before = Pt(18 if i <= 2 else 12)
        h.paragraph_format.space_after = Pt(8)

    try:
        code = doc.styles.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH)
    except ValueError:
        code = doc.styles["CodeBlock"]
    code.font.name = "Consolas"
    code.font.size = Pt(9.5)
    code.font.color.rgb = RGBColor(0xEE, 0xEE, 0xEE)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(10)
    code.paragraph_format.left_indent = Cm(0.3)


def shade_paragraph(paragraph, hex_color):
    pPr = paragraph._p.get_or_add_pPr()
    shd = pPr.makeelement(qn("w:shd"), {})
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def add_code_block(doc: Document, code_text: str) -> None:
    p = doc.add_paragraph(style="CodeBlock")
    shade_paragraph(p, "2B2B2B")
    lines = code_text.strip("\n").split("\n")
    for i, line in enumerate(lines):
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xEE, 0xEE, 0xEE)
        if i < len(lines) - 1:
            run.add_break()


def add_bullets(doc: Document, items, style="List Bullet"):
    for item in items:
        p = doc.add_paragraph(style=style)
        _add_rich_text(p, item)


def _add_rich_text(paragraph, text: str) -> None:
    """Supports **bold** and `code` inline markers."""
    i = 0
    while i < len(text):
        if text[i : i + 2] == "**":
            end = text.find("**", i + 2)
            if end != -1:
                r = paragraph.add_run(text[i + 2 : end])
                r.bold = True
                i = end + 2
                continue
        if text[i] == "`":
            end = text.find("`", i + 1)
            if end != -1:
                r = paragraph.add_run(text[i + 1 : end])
                r.font.name = "Consolas"
                r.font.size = Pt(9.8)
                i = end + 1
                continue
        # plain run until next special char
        next_special = min(
            [p for p in (text.find("**", i), text.find("`", i)) if p != -1] or [len(text)]
        )
        paragraph.add_run(text[i:next_special])
        i = next_special


def add_para(doc: Document, text: str, italic=False, color=None, size=None) -> None:
    p = doc.add_paragraph()
    _add_rich_text(p, text)
    if italic or color or size:
        for r in p.runs:
            if italic:
                r.italic = True
            if color:
                r.font.color.rgb = color
            if size:
                r.font.size = Pt(size)


def add_table(doc: Document, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=10.5)
        set_cell_shading(hdr[i], HEADER_FILL)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cell = cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            _add_rich_text(p, str(val))
            for run in p.runs:
                run.font.size = Pt(9.8)
            if r_idx % 2 == 1:
                set_cell_shading(cell, LIGHT_FILL)
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def add_step_heading(doc: Document, number: str, title: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    run = p.add_run(f"Step {number}  ")
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = ACCENT
    run2 = p.add_run(title)
    run2.bold = True
    run2.font.size = Pt(13)
    run2.font.color.rgb = NAVY


def add_page_break(doc: Document):
    doc.add_page_break()


def add_cover(doc: Document):
    for _ in range(4):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("AQI Tracking System")
    r.font.size = Pt(38)
    r.bold = True
    r.font.color.rgb = NAVY

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Project Walkthrough")
    r.font.size = Pt(22)
    r.font.color.rgb = ACCENT

    sub2 = doc.add_paragraph()
    sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub2.add_run("Open-source streaming + batch lakehouse for real-time Air Quality Index monitoring")
    r.font.size = Pt(12)
    r.italic = True
    r.font.color.rgb = GREY

    for _ in range(3):
        doc.add_paragraph()

    if ARCH_PNG.exists():
        doc.add_picture(str(ARCH_PNG), width=Cm(16.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    for _ in range(2):
        doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        "Kafka  ·  Flink  ·  MinIO  ·  Hive Metastore  ·  Trino  ·  Airflow  ·  Prometheus  ·  Alertmanager  ·  Grafana"
    )
    r.font.size = Pt(10.5)
    r.font.color.rgb = GREY


def add_toc(doc: Document):
    doc.add_heading("Contents", level=1)
    entries = [
        "1. Introduction & Purpose",
        "2. Architecture at a Glance",
        "3. Technology Stack",
        "4. Repository Layout",
        "5. Walkthrough: From Sensor Reading to Dashboard",
        "   5.1 Ingestion — Sensor Simulator → Kafka",
        "   5.2 Real-Time Processing — Flink Detection Job",
        "   5.3 Landing Zone — Raw Data Archival",
        "   5.4 Batch ETL — Airflow + Trino",
        "   5.5 Monitoring & Alerting",
        "   5.6 Visualization — Grafana Dashboards",
        "6. Running the Project",
        "7. Verifying the Deployment",
        "8. Testing & Quality Assurance",
        "9. Operating the System",
        "10. Production Hardening & Next Steps",
        "Appendix A — Environment Variables",
        "Appendix B — Service & Port Reference",
    ]
    for e in entries:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        indent = e.startswith("   ")
        run = p.add_run(e.strip())
        run.font.size = Pt(11 if not indent else 10.3)
        run.font.color.rgb = NAVY if not indent else GREY
        if indent:
            p.paragraph_format.left_indent = Cm(0.6)


# --------------------------------------------------------------------------- build
def main():
    doc = Document()
    build_styles(doc)

    add_cover(doc)
    add_page_break(doc)
    add_toc(doc)
    add_page_break(doc)

    # 1. Introduction
    doc.add_heading("1. Introduction & Purpose", level=1)
    add_para(
        doc,
        "The **AQI Tracking System** is a production-style, fully open-source reference platform for "
        "ingesting, processing, storing, querying, monitoring and visualizing Air Quality Index (AQI) "
        "data. It solves a problem most reference architectures split into two disconnected halves: "
        "**real-time visibility** (is the air unhealthy right now, and who should be paged) and "
        "**durable historical analytics** (trends, compliance reporting, BI). This project builds both "
        "against the *same* AQI calculation logic, so the live dashboard and the historical report never "
        "disagree about what \"AQI 160\" means.",
    )
    add_para(doc, "Who this is for:")
    add_bullets(
        doc,
        [
            "**Data / platform engineers** who want a runnable Kafka → Flink → Lakehouse → BI reference "
            "architecture to adapt for a real sensor fleet.",
            "**Environmental / public-health teams** who need live AQI dashboards and threshold "
            "alerting (email, Slack, SMS) without a SaaS platform.",
            "**Engineers learning the stack** — Kafka, Flink, Hive, Trino, Airflow, Prometheus and "
            "Grafana are all wired together and verified to run end-to-end, which is usually the hardest "
            "part of learning any one of them from documentation alone.",
        ],
    )

    # 2. Architecture
    doc.add_heading("2. Architecture at a Glance", level=1)
    add_para(
        doc,
        "The diagram below (generated from `docs/architecture.py`, kept in sync with the code) shows "
        "every service and how data moves between them. Solid arrows are data flow; dashed arrows are "
        "metrics or checkpoint state.",
    )
    if ARCH_PNG.exists():
        doc.add_picture(str(ARCH_PNG), width=Cm(17.0))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, "At a high level, the pipeline is:", size=10.5)
    add_code_block(
        doc,
        "sensor-simulator --(JSON, key=sensor_id)--> Kafka(aqi.raw)\n"
        "Kafka(aqi.raw) --> Flink(aqi-detection) --> Kafka(enriched / alerts / anomalies / dlq)\n"
        "Kafka(aqi.raw) --> Flink(aqi-landing-zone) --> MinIO(landing/aqi/raw/dt=/hour=)\n"
        "Kafka(enriched/alerts/anomalies) --> aqi-exporter --> Prometheus --> Grafana / Alertmanager\n"
        "Airflow --(Trino SQL)--> Trino <-> Hive Metastore <-> Postgres\n"
        "Trino --(read landing, write curated+analytics Parquet)--> MinIO(processed/)\n"
        "Alertmanager --> Mailpit (email) / webhook (Slack, SMS relay)",
    )

    # 3. Technology stack
    doc.add_heading("3. Technology Stack", level=1)
    add_table(
        doc,
        ["Layer", "Technology", "Role in this project"],
        [
            ["Ingestion", "Apache Kafka 3.8 (KRaft)", "`aqi.raw` topic, 12 partitions keyed by sensor_id"],
            ["Stream processing", "Apache Flink 1.18 (PyFlink)", "Per-reading AQI, threshold alerts, anomaly detection, landing-zone sink"],
            ["Object storage", "MinIO (S3 API)", "`landing` (raw JSON) and `processed` (curated Parquet) zones"],
            ["Catalog", "Hive Metastore 4.0 + Postgres", "Table & partition metadata for raw, curated, aggregated datasets"],
            ["Query engine", "Trino 455", "Ad-hoc SQL, ETL engine, BI / Grafana backend"],
            ["Orchestration", "Apache Airflow 2.10", "Daily curation, daily/weekly aggregation, partition sync, data-quality checks"],
            ["Monitoring", "Prometheus + Alertmanager", "Pipeline health & AQI threshold alerts via email, Slack, SMS"],
            ["Visualization", "Grafana 11", "Real-time AQI (Prometheus) + historical trends (Trino)"],
            ["Containerization", "Docker Compose v2", "~20 services, 5 custom-built images"],
            ["Testing", "pytest, DuckDB, sqlglot, Jinja2", "53 unit tests incl. SQL-vs-Python equivalence checks"],
        ],
        widths=[3.2, 4.8, 9.0],
    )

    # 4. Repo layout
    doc.add_heading("4. Repository Layout", level=1)
    add_code_block(
        doc,
        "aqi_common/              shared AQI + detection logic (used by Flink AND Airflow/Trino)\n"
        "generator/               sensor simulator (Kafka producer) + Dockerfile\n"
        "kafka/create-topics.sh   topic provisioning (partitions, retention)\n"
        "flink/                   PyFlink image, jobs/ (detection, landing zone), submit-jobs.sh\n"
        "hive/                    Hive Metastore image (Postgres JDBC + S3A) and hive-site.xml\n"
        "trino/catalog/           Trino Hive connector config (MinIO native S3)\n"
        "trino/sql/ddl.sql        schemas + tables for raw, curated, aggregated datasets\n"
        "airflow/                 Airflow image, DAGs and templated Trino SQL\n"
        "exporter/                Kafka -> Prometheus bridge for AQI business metrics\n"
        "monitoring/prometheus/   scrape config + alert rules\n"
        "monitoring/alertmanager/ routing (email/Slack/SMS), message templates\n"
        "grafana/                 provisioning (datasources, dashboards) + dashboard JSON\n"
        "minio/, postgres/        bucket/lifecycle and database bootstrap\n"
        "tests/                   unit tests (AQI math, detection, SQL templates)\n"
        "docs/                    architecture diagram + this walkthrough",
    )

    # 5. Walkthrough
    add_page_break(doc)
    doc.add_heading("5. Walkthrough: From Sensor Reading to Dashboard", level=1)
    add_para(
        doc,
        "This section follows one sensor reading through the entire system, in the order it actually "
        "happens.",
    )

    add_step_heading(doc, "5.1", "Ingestion — Sensor Simulator → Kafka")
    add_para(
        doc,
        "`generator/sensor_simulator.py` models a fleet of sensors across 8 cities with realistic "
        "diurnal traffic patterns, occasional multi-minute \"pollution episodes,\" and injected sensor "
        "faults (spikes, negative values, dropouts, stuck sensors). Each reading is published as JSON to "
        "the `aqi.raw` Kafka topic, keyed by `sensor_id` so a sensor's readings stay ordered within one "
        "partition.",
    )
    add_code_block(
        doc,
        '{"schema_version":1,"sensor_id":"delhi-0000","city":"Delhi","latitude":28.50,"longitude":77.31,\n'
        ' "pm25":64.17,"pm10":112.75,"no2":33.27,"o3":18.96,"co2":463.1,"firmware":"2.3.1",\n'
        ' "event_time":"2026-09-23T03:27:47.948Z"}',
    )
    add_para(
        doc,
        "`kafka/create-topics.sh` provisions five topics up front: `aqi.raw` (12 partitions, 7-day "
        "retention), `aqi.enriched` (12 partitions, 3-day retention), `aqi.alerts` / `aqi.anomalies` "
        "(3 partitions, 30-day retention for audit), and `aqi.dlq` for malformed messages.",
    )

    add_step_heading(doc, "5.2", "Real-Time Processing — Flink Detection Job")
    add_para(doc, "`flink/jobs/aqi_detection_job.py` consumes `aqi.raw` and, per reading:")
    add_bullets(
        doc,
        [
            "**Parses & validates** the JSON (`aqi_common.detection.parse_reading`) — malformed "
            "messages are routed to `aqi.dlq` instead of crashing the job.",
            "**Computes AQI** (`aqi_common.aqi.compute_aqi`) using the 2024 EPA breakpoint tables for "
            "PM2.5 / PM10 / NO2 / O3, taking the maximum sub-index and recording the dominant pollutant. "
            "CO2 gets its own indoor-air-quality level.",
            "**Keeps per-sensor keyed state**, checkpointed to MinIO every 30 seconds:",
        ],
    )
    add_bullets(
        doc,
        [
            "**`ThresholdAlerter`** — severity escalates warning → critical → emergency as AQI/CO2 cross "
            "fixed limits (AQI ≥151/≥201/≥301; CO2 ≥1500/≥2000/≥5000 ppm); re-fires after a 15-minute "
            "cooldown (event-time based); emits `RESOLVED` on recovery. Output: `aqi.alerts`.",
            "**`AnomalyDetector`** — an EWMA mean/variance per pollutant per sensor flags z-score "
            "spikes/drops (|z| ≥ 4 after a 20-reading warm-up), physically impossible values, and "
            "\"stuck\" sensors repeating the same value 30 times in a row. Output: `aqi.anomalies`.",
        ],
        style="List Bullet 2",
    )
    add_para(
        doc,
        "Every reading, enriched with its AQI, is written to `aqi.enriched` — the topic the live "
        "dashboard reads from.",
    )

    add_step_heading(doc, "5.3", "Landing Zone — Raw Data Archival")
    add_para(
        doc,
        "`flink/jobs/landing_zone_job.py` is a separate PyFlink SQL job (independently scalable from "
        "detection) that archives every raw reading as newline-delimited JSON to "
        "`s3a://landing/aqi/raw/dt=YYYY-MM-DD/hour=HH/` in MinIO. Files are committed exactly once, tied "
        "to Flink's checkpoint cycle, so a crash never produces duplicate or partial files.",
    )

    add_step_heading(doc, "5.4", "Batch ETL — Airflow + Trino")
    add_para(
        doc,
        "`airflow/dags/aqi_daily_aggregation.py` runs nightly at 00:30 UTC and executes, entirely in "
        "Trino SQL:",
    )
    add_bullets(
        doc,
        [
            "**Sync partitions** — register any new landing-zone folders in the Hive Metastore.",
            "**Wait for data** — a sensor that ensures the day's raw data has actually landed.",
            "**Build `curated.readings`** — de-duplicate, drop physically implausible values, compute "
            "AQI (via Trino SQL generated from the *same* `aqi_common.aqi` module used by Flink — see "
            "the Testing section), write partitioned Parquet.",
            "**Data-quality check** — row counts, duplicate keys, null-AQI ratio, timestamp bounds "
            "(`SQLCheckOperator`).",
            "**Aggregate** — `daily_sensor_aqi`, `daily_city_aqi`, then `weekly_city_aqi`.",
            "**`ANALYZE`** — refresh table statistics for the query planner.",
        ],
    )
    add_para(
        doc,
        "Every step overwrites its own partition, so re-running or backfilling any day is always safe. "
        "A second DAG, `aqi_landing_partition_sync`, registers new raw partitions every 15 minutes so "
        "ad-hoc SQL sees near-real-time data without waiting for the nightly run.",
    )

    add_step_heading(doc, "5.5", "Monitoring & Alerting")
    add_para(
        doc,
        "`exporter/aqi_exporter.py` bridges the Kafka business topics (`aqi.enriched`, `aqi.alerts`, "
        "`aqi.anomalies`) into Prometheus metrics — latest AQI per sensor, active alerts, anomaly rates. "
        "Prometheus rules (`monitoring/prometheus/rules/`) split into two groups:",
    )
    add_bullets(
        doc,
        [
            "**Pipeline health** — Kafka consumer lag, Flink job state, checkpoint failures, dead-letter "
            "queue growth, MinIO capacity.",
            "**AQI business alerts** — city/sensor thresholds (`AQIUnhealthy`, `AQIVeryUnhealthy`, "
            "`AQIHazardous`), sensor silence, anomaly bursts.",
        ],
    )
    add_para(
        doc,
        "Alertmanager routes each alert by its `category`/`severity`/`team` labels to email (via "
        "Mailpit locally), Slack, and an SMS webhook, with inhibition rules so a Hazardous alert "
        "suppresses the redundant lower-severity alert for the same city.",
    )

    add_step_heading(doc, "5.6", "Visualization — Grafana Dashboards")
    add_para(
        doc,
        "Two dashboards are auto-provisioned in Grafana's **AQI** folder:",
    )
    add_bullets(
        doc,
        [
            "**AQI Overview** — current worst/median AQI, AQI by city, pollutant explorer, top sensors, "
            "streaming alert/anomaly rates, firing Alertmanager alerts, an AQI category timeline, and "
            "historical panels (daily average AQI, % unhealthy readings, weekly summary) sourced from "
            "Trino.",
            "**AQI Pipeline Health** — Flink job/TaskManager status, Kafka consumer lag per "
            "group/partition, topic throughput, DLQ rate, checkpoint duration, restarts, backpressure, "
            "and MinIO capacity.",
        ],
    )

    # 6. Running
    add_page_break(doc)
    doc.add_heading("6. Running the Project", level=1)
    add_para(doc, "**Prerequisites:** Docker Desktop / Engine 24+ with Compose v2, ~12 GB RAM allocated, ~10 GB free disk.")
    add_para(doc, "**Start the full stack:**")
    add_code_block(
        doc,
        "cd AQI\n"
        "cp .env.example .env          # optional — every value has a working local default\n"
        "docker compose up -d --build\n"
        "docker compose ps",
    )
    add_para(
        doc,
        "The first build downloads PyFlink, Hive JDBC/S3A and Airflow provider dependencies (several "
        "minutes). Startup order is automatic via Compose's `depends_on` + healthchecks: "
        "`postgres, kafka, minio` → `kafka-init, minio-init` → `jobmanager, taskmanager` → "
        "`flink-job-submitter` → `hive-metastore` → `trino` → `trino-init` → `airflow-*`. The one-shot "
        "init containers should show `Exited (0)` when done; everything else should be `running`/`healthy`.",
    )
    add_para(doc, "**Backfill history so the batch dashboards have data:**")
    add_code_block(
        doc,
        "docker compose run --rm sensor-simulator \\\n"
        "  --bootstrap=kafka:9092 --sensors=40 --backfill-hours=48 --backfill-step=300 --exit-after-backfill\n\n"
        "# then trigger the daily DAG for a specific day in the Airflow UI or:\n"
        'docker compose exec airflow-scheduler airflow dags trigger aqi_daily_aggregation \\\n'
        "  -c '{\"target_date\": \"2026-09-24\"}'",
    )

    # 7. Verifying
    doc.add_heading("7. Verifying the Deployment", level=1)
    add_bullets(
        doc,
        [
            "**Flink jobs running:** `curl localhost:8081/jobs/overview` should list `aqi-detection` "
            "and `aqi-landing-zone` both `RUNNING`.",
            "**Messages flowing:** "
            "`docker compose exec kafka /opt/kafka/bin/kafka-console-consumer.sh --bootstrap-server "
            "kafka:9092 --topic aqi.alerts --max-messages 5`",
            "**Files landing in MinIO:** open the console at `localhost:9001` and browse the "
            "`landing` bucket, or use `mc ls -r` inside the `minio-init` container.",
            "**Metrics exposed:** `curl localhost:8000/metrics` (aqi-exporter) should show `aqi_index` "
            "series per sensor.",
            "**Dashboards populated:** open Grafana at `localhost:3000` (admin/admin) → **AQI** folder.",
        ],
    )
    add_para(
        doc,
        "This exact sequence was run against a live deployment during development: both Flink jobs "
        "reached `RUNNING`, `aqi.enriched` produced correct `aqi_index` values in the exporter, a "
        "48-hour backfill (11,520 messages) and a full `aqi_daily_aggregation` run completed with all 8 "
        "tasks green, producing 5,760 curated rows and city aggregates matching the simulated pollution "
        "episodes (e.g., Delhi: average AQI 193, 94.7% of readings Unhealthy), and a live `AQIHazardous` "
        "alert was delivered to the Mailpit inbox end-to-end.",
        italic=True,
        color=GREY,
    )

    # 8. Testing
    doc.add_heading("8. Testing & Quality Assurance", level=1)
    add_para(doc, "Run with:")
    add_code_block(doc, "pip install -r requirements-dev.txt\npython -m pytest tests -q     # 53 passed")
    add_table(
        doc,
        ["Test file", "What it validates"],
        [
            [
                "`tests/test_aqi.py`",
                "EPA AQI breakpoint math (truncation, boundaries, categories), and — critically — that "
                "the **generated Trino SQL produces identical results to the Python implementation** "
                "(verified by executing the SQL in DuckDB against thousands of random values per "
                "pollutant). This is what guarantees the real-time and batch AQI numbers never diverge.",
            ],
            [
                "`tests/test_detection.py`",
                "Reading parsing/validation, the `AnomalyDetector` state machine (warm-up, spikes, "
                "stuck sensors), and the `ThresholdAlerter` state machine (escalation, cooldown, "
                "resolution, JSON-serializable state for Flink checkpointing).",
            ],
            [
                "`tests/test_sql_templates.py`",
                "Every templated Airflow SQL file renders to valid, parseable Trino SQL (via sqlglot), "
                "including parameter overrides, and that the DDL script parses correctly.",
            ],
        ],
        widths=[4.5, 12.5],
    )
    add_para(
        doc,
        "Two real defects were found and fixed via live end-to-end testing rather than unit tests alone: "
        "a PyFlink packaging issue (the shared `aqi_common` module needed to be shipped as a zip, not a "
        "directory, for Python UDF workers to import it) and a Hive Metastore 4 / Trino incompatibility "
        "with `ANALYZE` on timestamp columns. Both are documented in the README troubleshooting section.",
    )

    # 9. Operating
    add_page_break(doc)
    doc.add_heading("9. Operating the System", level=1)
    add_table(
        doc,
        ["Task", "Command"],
        [
            ["Scale Flink processing", "`docker compose up -d --scale taskmanager=3`, raise `DETECTION_PARALLELISM`, resubmit"],
            ["Resubmit Flink jobs (idempotent)", "`docker compose run --rm flink-job-submitter`"],
            ["Reload Prometheus rules", "`curl -X POST localhost:9090/-/reload`"],
            ["Reprocess a day", "Trigger `aqi_daily_aggregation` with `target_date`; partitions overwrite safely"],
            ["Tear down (keep data)", "`docker compose down`"],
            ["Tear down (delete all data)", "`docker compose down -v`"],
        ],
        widths=[6.0, 11.0],
    )

    # 10. Next steps
    doc.add_heading("10. Production Hardening & Next Steps", level=1)
    add_bullets(
        doc,
        [
            "**Kafka:** 3+ brokers, dedicated KRaft controller quorum, replication factor 3, TLS/SASL, "
            "and a schema registry (Avro/Protobuf) for the sensor contract.",
            "**Flink:** run on Kubernetes via the Flink Kubernetes Operator, JobManager HA, RocksDB "
            "state backend with incremental checkpoints, periodic savepoints.",
            "**MinIO:** distributed erasure-coded mode (4+ nodes), TLS, per-zone bucket policies, "
            "scoped service accounts instead of the root user.",
            "**Trino / Hive Metastore:** HA Postgres, multiple Trino workers, OAuth2/LDAP "
            "authentication, and access control (file-based rules or Apache Ranger).",
            "**Table format:** consider migrating the curated zone to Apache Iceberg (same Hive "
            "Metastore, Trino's `iceberg` connector) for ACID `MERGE`, schema evolution and time travel.",
            "**Airflow:** CeleryExecutor/KubernetesExecutor, secrets in a real secrets backend "
            "(Vault, AWS Secrets Manager), `on_failure_callback` notifications.",
            "**CI/CD:** no pipeline is configured yet — add one to run `pytest`, "
            "`docker compose config --quiet`, and build/smoke-test the custom images on every PR.",
            "**Feature ideas:** a thin read-only public API over Trino, geospatial interpolation "
            "between sensors for AQI heatmaps, and per-subscriber alerting built on the existing "
            "`aqi.alerts` topic.",
        ],
    )

    # Appendix A
    add_page_break(doc)
    doc.add_heading("Appendix A — Key Environment Variables", level=1)
    add_para(doc, "Set in `.env` (copy from `.env.example`); every value has a working local default.")
    add_table(
        doc,
        ["Variable", "Purpose"],
        [
            ["MINIO_ROOT_USER / MINIO_ROOT_PASSWORD", "MinIO credentials"],
            ["POSTGRES_USER / POSTGRES_PASSWORD", "Shared Postgres instance (Metastore + Airflow DBs)"],
            ["AIRFLOW_ADMIN_USER / AIRFLOW_ADMIN_PASSWORD", "Airflow UI login"],
            ["ALERT_EMAIL_TO / SMTP_*", "Where and how Alertmanager sends email"],
            ["SLACK_WEBHOOK_URL / SLACK_CHANNEL_*", "Slack alert routing"],
            ["SMS_WEBHOOK_URL", "Any relay accepting Alertmanager's webhook JSON (Twilio, SNS, ...)"],
            ["AQI_RAW_PARTITIONS", "Partition count for `aqi.raw` (scale with consumer parallelism)"],
            ["ALERT_COOLDOWN_MINUTES / ANOMALY_Z_THRESHOLD", "Streaming detection tuning"],
            ["SIM_SENSORS / SIM_INTERVAL_SECONDS / SIM_ANOMALY_RATE", "Sensor simulator load & fault rate"],
            ["LANDING_RETENTION_DAYS / PROMETHEUS_RETENTION", "Storage retention"],
        ],
        widths=[6.5, 10.5],
    )

    # Appendix B
    doc.add_heading("Appendix B — Service & Port Reference", level=1)
    add_table(
        doc,
        ["Service", "URL", "Credentials"],
        [
            ["Grafana", "http://localhost:3000", "admin / admin"],
            ["Flink UI", "http://localhost:8081", "—"],
            ["Trino UI", "http://localhost:8080", "any user"],
            ["Airflow", "http://localhost:8088", "admin / admin"],
            ["MinIO console", "http://localhost:9001", "aqiadmin / aqiadmin-secret"],
            ["Prometheus", "http://localhost:9090", "—"],
            ["Alertmanager", "http://localhost:9093", "—"],
            ["Mailpit (alert emails)", "http://localhost:8025", "—"],
            ["Kafka (from host)", "localhost:9094", "—"],
            ["AQI exporter metrics", "http://localhost:8000/metrics", "—"],
        ],
        widths=[4.5, 7.0, 5.5],
    )

    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer_p.add_run("Generated from the AQI Tracking System repository — see README.md for full reference documentation.")
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = GREY

    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
